# -*- coding: utf-8 -*-

import os
import pythoncom
from win32com.client import Dispatch, gencache
import openpyxl


def KompasAPI():
    global KompasAPI7
    global iApplication
    global ksConstants

    KompasAPI7 = gencache.EnsureModule(
        '{69AC2981-37C0-4379-84FD-5DD2F3C0A520}', 0, 1, 0)
    iApplication = Dispatch('Kompas.Application.7')
    ksConstants = gencache.EnsureModule(
        '{75C9F5D0-B5B8-4526-8681-9903C567D2ED}', 0, 1, 0).constants


def read_excel_data(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [cell.value for cell in ws[1]]
    wb.close()
    return headers, rows


def create_table():
    global KompasAPI7, iApplication, ksConstants

    excel_path = os.path.join(os.getcwd(), 'loodsman_report_ko_101388.xlsx')
    if not os.path.exists(excel_path):
        print("Файл Excel не найден: {}".format(excel_path))
        return

    headers, data = read_excel_data(excel_path)
    print("Прочитано строк из Excel: {}".format(len(data)))
    print("Заголовки: {}".format(headers))

    iKompasDocument = iApplication.ActiveDocument
    if not iKompasDocument:
        print("Активный документ не найден")
        return

    iKompasDocument2D = KompasAPI7.IKompasDocument2D(iKompasDocument)
    iViewsAndLayersManager = iKompasDocument2D.ViewsAndLayersManager
    iViews = iViewsAndLayersManager.Views
    iView = iViews.ActiveView

    # получаем контейнер символов и коллекцию таблиц
    iSymbols2DContainer = KompasAPI7.ISymbols2DContainer(iView)
    iDrawingTables = KompasAPI7.IDrawingTables(iSymbols2DContainer.DrawingTables)

    rows_count = len(data) + 1
    columns_count = 5
    row_height = 8.0
    column_width = 35.0

    if iDrawingTables is None:
        print("Не удалось получить коллекцию таблиц")
        return

    # создаём таблицу
    title_pos = getattr(ksConstants, 'ksTableTitleLayoutNone',
                        getattr(ksConstants, 'ksTableTileLayoutNone', 0))
    iDrawingTable = KompasAPI7.IDrawingTable(iDrawingTables.Add(
        rows_count, columns_count, row_height, column_width,
        title_pos
    ))

    # устанавливаем позицию таблицы на виде (мм)
    iDrawingTable.X = 100.0
    iDrawingTable.Y = 200.0

    # получаем интерфейс сетки таблицы
    iTable = KompasAPI7.ITable(iDrawingTable)

    # заполняем заголовки
    header_values = [
        'Наименование детали',
        'Количество на изделие',
        'Обозначение детали',
        'Наименование материала',
        'Марка материала'
    ]
    for col_idx, header_text in enumerate(header_values):
        cell_id = 0 * columns_count + col_idx + 1
        iTableCell = iTable.CellById(cell_id)
        iText = KompasAPI7.IText(iTableCell.Text)
        iText.Str = header_text

    # заполняем данные
    for row_idx, row_data in enumerate(data):
        for col_idx in range(columns_count):
            cell_id = (row_idx + 1) * columns_count + col_idx + 1
            iTableCell = iTable.CellById(cell_id)
            iText = KompasAPI7.IText(iTableCell.Text)
            value = row_data[col_idx]
            if value is None:
                value = ''
            else:
                value = str(value)
            iText.Str = value

    # обновляем объект таблицы, чтобы изменения вступили в силу
    iDrawingTable.Update()

    print("Таблица успешно создана: {} строк x {} столбцов".format(rows_count, columns_count))


KompasAPI()
create_table()
